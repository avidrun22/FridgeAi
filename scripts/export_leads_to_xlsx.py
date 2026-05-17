#!/usr/bin/env python3
"""
Export ok2eat leads to an .xlsx with two tabs.

Pulls every row from `business_waitlist` (B2B form) and `android_waitlist`
(homepage + /quiz Android signup) via the Supabase REST API, writes a
spreadsheet with one tab per source, and saves to a local path.

Designed to be re-run as an automation:
  - reads credentials from .appstoreconnect/telegram_config.json (or env vars)
  - writes to /tmp/leads.xlsx by default, override with --out
  - exit 0 with the file path printed to stdout on success
  - non-zero with a stderr message on failure

A companion Cowork scheduled task picks up the output file and uploads it
to Google Drive via the Drive MCP, replacing the previous "leads" sheet.

Usage:
  python3 scripts/export_leads_to_xlsx.py                  # writes /tmp/leads.xlsx
  python3 scripts/export_leads_to_xlsx.py --out ./leads.xlsx
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path
from typing import Any

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH  = PROJECT_ROOT / ".appstoreconnect" / "telegram_config.json"

TABLES = [
    # (Supabase table, sheet tab name, columns to surface in order)
    (
        "business_waitlist",
        "B2B Leads",
        # Pulled dynamically — we include everything the row has so we don't
        # silently drop a column when the schema grows.
        None,
    ),
    (
        "android_waitlist",
        "Android Waitlist",
        None,
    ),
]


def load_credentials() -> tuple[str, str]:
    """Read Supabase URL + service-role key from env or telegram_config.json."""
    url = os.environ.get("SUPABASE_URL")
    key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if url and key:
        return url.rstrip("/"), key
    if not CONFIG_PATH.exists():
        sys.exit(f"[export_leads] config not found at {CONFIG_PATH} and no env vars set")
    cfg = json.loads(CONFIG_PATH.read_text())
    url = cfg.get("SUPABASE_URL") or cfg.get("supabase_url")
    key = cfg.get("SUPABASE_SERVICE_ROLE_KEY") or cfg.get("supabase_service_role_key")
    if not url or not key:
        sys.exit("[export_leads] supabase_url or supabase_service_role_key missing from config")
    return url.rstrip("/"), key


def fetch_table(supabase_url: str, service_key: str, table: str) -> list[dict[str, Any]]:
    """SELECT * via PostgREST. Newest rows first; cap at 1000."""
    endpoint = f"{supabase_url}/rest/v1/{table}"
    # Order by the first plausible timestamp column. Both tables happen to have
    # `created_at` (business) or `signed_up_at` (android) — try both, fall back
    # to no ordering if neither exists (PostgREST 400s on bad columns).
    for order_col in ("signed_up_at.desc", "created_at.desc", None):
        params: dict[str, str] = {"select": "*", "limit": "1000"}
        if order_col:
            params["order"] = order_col
        resp = requests.get(
            endpoint,
            params=params,
            headers={
                "Authorization": f"Bearer {service_key}",
                "apikey": service_key,
                "Accept": "application/json",
            },
            timeout=30,
        )
        if resp.status_code == 400 and order_col:
            # bad order column — try next fallback
            continue
        resp.raise_for_status()
        return resp.json()
    return []


def write_workbook(out_path: Path, all_rows: list[tuple[str, list[dict[str, Any]]]]) -> None:
    wb = Workbook()
    # Remove the default sheet and create our own (otherwise we end up with an
    # empty "Sheet" tab alongside ours).
    default = wb.active
    wb.remove(default)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="3E721D", end_color="3E721D", fill_type="solid")
    header_align = Alignment(horizontal="left", vertical="center", wrap_text=False)

    for sheet_name, rows in all_rows:
        ws = wb.create_sheet(title=sheet_name)
        if not rows:
            ws["A1"] = f"No rows in {sheet_name} yet."
            ws["A1"].font = Font(italic=True, color="808080")
            continue

        # Union of all keys across rows preserves the natural column order of
        # the first row, then appends any extras from later rows.
        columns: list[str] = []
        seen: set[str] = set()
        for row in rows:
            for k in row.keys():
                if k not in seen:
                    columns.append(k)
                    seen.add(k)

        # Header row
        for col_idx, key in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=key)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # Data rows
        for row_idx, row in enumerate(rows, start=2):
            for col_idx, key in enumerate(columns, start=1):
                value = row.get(key)
                # Flatten JSON-ish values to strings so openpyxl doesn't choke.
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, separators=(",", ":"), default=str)
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Auto-size each column to its max content width (capped so we don't
        # blow up on long notes fields).
        for col_idx, key in enumerate(columns, start=1):
            max_len = len(key)
            for row in rows:
                v = row.get(key)
                if v is None:
                    continue
                s = str(v)
                if len(s) > max_len:
                    max_len = len(s)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 50)

        ws.freeze_panes = "A2"  # keep header visible when scrolling

    wb.save(out_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Export ok2eat leads to xlsx")
    parser.add_argument("--out", default="/tmp/leads.xlsx", help="Output xlsx path")
    args = parser.parse_args()

    supabase_url, service_key = load_credentials()

    all_rows: list[tuple[str, list[dict[str, Any]]]] = []
    totals: list[str] = []
    for table, sheet_name, _ in TABLES:
        rows = fetch_table(supabase_url, service_key, table)
        all_rows.append((sheet_name, rows))
        totals.append(f"{sheet_name}: {len(rows)}")
        print(f"[export_leads] fetched {len(rows)} rows from {table}", file=sys.stderr)

    out_path = Path(args.out).expanduser().resolve()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    write_workbook(out_path, all_rows)

    print(str(out_path))
    print(f"[export_leads] wrote {out_path} ({' · '.join(totals)})", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
